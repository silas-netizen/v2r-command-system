"""대량 원고 생성 파이프라인(brand_queue + bulk_generate) 테스트."""

from __future__ import annotations

import json
from pathlib import Path

import pytest

from tests.test_brand_writer import COMMENTS, FakeLLM, BODY, KEYWORD
from tests.test_engine import make_runtime
from v2r.command.parser import parse_korean_command
from v2r.content import brand_queue, bulk_generate
from v2r.engine import worker


def _patch_pool(monkeypatch, items):
    monkeypatch.setattr(
        "v2r.sources.keyword_list.load_pushed_keywords",
        lambda brand, cfg=None, xlsx_path=None, limit=0, conn=None: items,
    )


def test_refill_inserts_pending_rows(tmp_path, monkeypatch):
    rt = make_runtime(tmp_path)
    rt.settings.repo_root = tmp_path
    _patch_pool(
        monkeypatch,
        [{"keyword": "비타민C", "cafe": ""}, {"keyword": "이노시톨", "cafe": ""}],
    )
    added = brand_queue.refill(rt, "우아덤", 2)
    assert added == 2
    rows = brand_queue.pending(rt, "우아덤")
    assert {r["keyword"] for r in rows} == {"비타민C", "이노시톨"}
    assert all(r["mtype"] == "질문형" for r in rows)
    rt.close()


def test_refill_skips_ready_or_published(tmp_path, monkeypatch):
    rt = make_runtime(tmp_path)
    rt.settings.repo_root = tmp_path
    _patch_pool(monkeypatch, [{"keyword": "비타민C", "cafe": ""}])
    brand_queue.refill(rt, "우아덤", 1)
    row = brand_queue.pending(rt, "우아덤")[0]
    brand_queue.mark(rt, row["id"], "ready", manuscript_path="x.json")
    added = brand_queue.refill(rt, "우아덤", 1)
    assert added == 0
    assert brand_queue.pending(rt, "우아덤") == []
    rt.close()


def test_refill_alternates_mtype_for_patsuni(tmp_path, monkeypatch):
    rt = make_runtime(tmp_path)
    rt.settings.repo_root = tmp_path
    _patch_pool(
        monkeypatch,
        [{"keyword": f"키워드{i}", "cafe": ""} for i in range(4)],
    )
    brand_queue.refill(rt, "팥순이", 4)
    rows = brand_queue.pending(rt, "팥순이")
    mtypes = [r["mtype"] for r in rows]
    assert mtypes.count("질문형") == 2 and mtypes.count("후기형") == 2
    rt.close()


def test_refill_orders_by_volume_then_relevance(tmp_path, monkeypatch):
    rt = make_runtime(tmp_path)
    rt.settings.repo_root = tmp_path
    _patch_pool(
        monkeypatch,
        [
            {"keyword": "낮음", "cafe": ""},
            {"keyword": "높음", "cafe": ""},
            {"keyword": "무관측", "cafe": ""},
        ],
    )
    import sqlite3

    kdir = tmp_path / "data" / "keywords"
    kdir.mkdir(parents=True, exist_ok=True)
    conn = sqlite3.connect(str(kdir / "우아덤.sqlite"))
    conn.execute(
        "CREATE TABLE keywords (keyword TEXT, pc INT, mobile INT, total INT,"
        " source_seed TEXT, depth INT, relevance INT, collected_at TEXT)"
    )
    conn.execute("INSERT INTO keywords VALUES ('낮음',10,10,20,'',1,0,'')")
    conn.execute("INSERT INTO keywords VALUES ('높음',100,100,200,'',1,5,'')")
    conn.commit()
    conn.close()

    brand_queue.refill(rt, "우아덤", 3)
    rows = brand_queue.pending(rt, "우아덤")
    assert [r["keyword"] for r in rows] == ["높음", "낮음", "무관측"]
    rt.close()


def test_bulk_generate_for_brand_saves_to_queue_folder(tmp_path, monkeypatch):
    rt = make_runtime(tmp_path)
    rt.settings.repo_root = tmp_path
    rt._llm = FakeLLM()
    rt._llm_ready = True
    _patch_pool(monkeypatch, [{"keyword": KEYWORD, "cafe": "씨씨앙"}])

    out = bulk_generate.generate_for_brand(rt, "우아덤", 1)
    assert out["ok"] is True
    assert out["generated"] == 1
    saved = (
        tmp_path / "warehouse" / "manuscripts" / "brand-queue" / "우아덤" / f"{KEYWORD}.json"
    )
    assert saved.exists()
    rows = list(rt.conn.execute("SELECT status FROM brand_queue WHERE brand='우아덤'"))
    assert rows and rows[0]["status"] == "ready"
    rt.close()


def test_bulk_generate_dedups_across_runs(tmp_path, monkeypatch):
    rt = make_runtime(tmp_path)
    rt.settings.repo_root = tmp_path
    rt._llm = FakeLLM()
    rt._llm_ready = True
    _patch_pool(monkeypatch, [{"keyword": KEYWORD, "cafe": "씨씨앙"}])

    first = bulk_generate.generate_for_brand(rt, "우아덤", 1)
    assert first["generated"] == 1
    second = bulk_generate.generate_for_brand(rt, "우아덤", 1)
    assert second["generated"] == 0
    rt.close()


def test_bulk_generate_no_paid_fallback_on_plan_limit(tmp_path, monkeypatch):
    """요금제 한도(PlanLimit)에 걸리면 유료로 넘어가지 않고 대기 상태로 멈춘다."""
    from v2r.llm.plan_backend import PlanLimit

    class LimitedLLM(FakeLLM):
        def __init__(self):
            super().__init__()
            self.force_backend = ""

        def plan_locked_until(self):
            return "2999-01-01T00:00:00+09:00"

    rt = make_runtime(tmp_path)
    rt.settings.repo_root = tmp_path
    rt._llm = LimitedLLM()
    rt._llm_ready = True
    _patch_pool(
        monkeypatch,
        [{"keyword": "A", "cafe": ""}, {"keyword": "B", "cafe": ""}],
    )

    monkeypatch.setattr(
        "v2r.content.bulk_generate.load_yaml",
        lambda name: {"bulk": {"allow_paid_fallback": False, "brands": ["우아덤"]}}
        if name == "bulk"
        else {},
    )

    out = bulk_generate.generate_for_brand(rt, "우아덤", 2)
    assert out["generated"] == 0
    assert out["waiting_plan_limit"] is True
    rows = list(rt.conn.execute("SELECT status FROM brand_queue WHERE brand='우아덤'"))
    assert all(r["status"] == "pending" for r in rows)
    rt.close()


def test_status_reports_counts(tmp_path, monkeypatch):
    rt = make_runtime(tmp_path)
    rt.settings.repo_root = tmp_path
    _patch_pool(monkeypatch, [{"keyword": "A", "cafe": ""}])
    brand_queue.refill(rt, "우아덤", 1)
    out = bulk_generate.status(rt)
    assert out["ok"] is True
    assert out["pending"] == 1
    rt.close()


def test_status_reports_by_target_and_room_to_max_ready(tmp_path, monkeypatch):
    """target(v2r/vpc)별 현황과 재고 상한 여유(사용자 지시 갱신 2026-09-23 01:35)."""
    rt = make_runtime(tmp_path)
    rt.settings.repo_root = tmp_path
    cfg = {"bulk": {"brands": ["우아덤"], "max_ready": 100, "v2r_daily_cap": 2}}
    monkeypatch.setattr(
        "v2r.content.bulk_generate.load_yaml",
        lambda name: cfg if name == "bulk" else {},
    )
    monkeypatch.setattr(
        "v2r.content.brand_queue.load_yaml",
        lambda name: cfg if name == "bulk" else {},
    )
    _patch_pool(
        monkeypatch,
        [{"keyword": f"키워드{i}", "cafe": ""} for i in range(3)],
    )
    brand_queue.refill(rt, "우아덤", 3)  # v2r_daily_cap=2 → 2건 v2r, 1건 vpc
    out = bulk_generate.status(rt)
    assert out["max_ready"] == 100
    assert out["room_to_max_ready"] == 100  # 아직 ready 0건
    assert out["by_target"]["v2r"]["pending"] == 2
    assert out["by_target"]["vpc"]["pending"] == 1
    rt.close()


def test_refill_auto_assigns_target_by_daily_cap(tmp_path, monkeypatch):
    """검색량 상위부터 하루 v2r_daily_cap까지는 v2r, 그다음은 vpc(사용자 지시 2026-09-23)."""
    rt = make_runtime(tmp_path)
    rt.settings.repo_root = tmp_path
    monkeypatch.setattr(
        "v2r.content.brand_queue.load_yaml",
        lambda name: {"bulk": {"v2r_daily_cap": 2}} if name == "bulk" else {},
    )
    _patch_pool(
        monkeypatch,
        [{"keyword": f"키워드{i}", "cafe": ""} for i in range(4)],
    )
    brand_queue.refill(rt, "우아덤", 4)
    rows = brand_queue.pending(rt, "우아덤")
    targets = [r["target"] for r in rows]
    assert targets.count("v2r") == 2 and targets.count("vpc") == 2
    rt.close()


def test_refill_target_override_forces_vpc(tmp_path, monkeypatch):
    rt = make_runtime(tmp_path)
    rt.settings.repo_root = tmp_path
    _patch_pool(monkeypatch, [{"keyword": "A", "cafe": ""}])
    brand_queue.refill(rt, "우아덤", 1, target="vpc")
    row = brand_queue.pending(rt, "우아덤")[0]
    assert row["target"] == "vpc"
    rt.close()


def test_generate_for_brand_stops_at_max_ready(tmp_path, monkeypatch):
    """재고(ready)가 max_ready에 닿으면 더 채우지 않는다(사용자 지시 갱신 2026-09-23)."""
    rt = make_runtime(tmp_path)
    rt.settings.repo_root = tmp_path
    rt._llm = FakeLLM()
    rt._llm_ready = True
    _patch_pool(monkeypatch, [{"keyword": KEYWORD, "cafe": "씨씨앙"}])
    monkeypatch.setattr(
        "v2r.content.bulk_generate.load_yaml",
        lambda name: {"bulk": {"allow_paid_fallback": False, "max_ready": 1}}
        if name == "bulk"
        else {},
    )
    out = bulk_generate.generate_for_brand(rt, "우아덤", 1)
    assert out["generated"] == 1
    out2 = bulk_generate.generate_for_brand(rt, "우아덤", 1)
    assert out2["generated"] == 0
    assert out2.get("stock_full") is True
    rt.close()


def test_export_vpc_writes_both_files_and_marks_exported(tmp_path, monkeypatch):
    """인수인계 형식(2026-09-23): 브랜드별 xlsx + 통합 댓글 프로그램용 xlsx/csv."""
    rt = make_runtime(tmp_path)
    rt.settings.repo_root = tmp_path
    rt._llm = FakeLLM()
    rt._llm_ready = True
    _patch_pool(monkeypatch, [{"keyword": KEYWORD, "cafe": "씨씨앙"}])
    monkeypatch.setattr(
        "v2r.content.bulk_generate.load_yaml",
        lambda name: {"bulk": {"allow_paid_fallback": False}} if name == "bulk" else {},
    )
    brand_queue.refill(rt, "우아덤", 1, target="vpc")
    out = bulk_generate.generate_for_brand(rt, "우아덤", 1, target="vpc")
    assert out["generated"] == 1

    export_out = bulk_generate.export_vpc(rt)
    assert export_out["exported"] == 1

    import openpyxl

    # 1) 브랜드별 xlsx — 시트 이름 = 키워드, 열 = 키워드/본문/../작성계정/원고유형/완료 링크
    brand_path = Path(export_out["brand_files"]["우아덤"])
    assert brand_path.exists()
    wb1 = openpyxl.load_workbook(str(brand_path))
    assert KEYWORD in wb1.sheetnames
    ws1 = wb1[KEYWORD]
    header1 = [c.value for c in next(ws1.iter_rows(max_row=1))]
    assert header1 == ["키워드", "본문", None, "작성계정", "원고유형", "완료 링크"]
    data_row = [c.value for c in list(ws1.iter_rows(min_row=2, max_row=2))[0]]
    assert data_row[0] == KEYWORD
    assert "댓글1:" in data_row[1] and "2.1:" in data_row[1] and "2.2:" in data_row[1]

    # 2) 통합 댓글 프로그램용 — 게시글마다 시트(완료 링크 없으니 시트명=키워드),
    #    열 3개 고정, 계정 배정표대로 채워진다
    prog_path = Path(export_out["comment_program_xlsx"])
    assert prog_path.exists()
    wb2 = openpyxl.load_workbook(str(prog_path))
    ws2 = wb2[KEYWORD]
    header2 = [c.value for c in next(ws2.iter_rows(max_row=1))]
    assert header2 == ["작성자 구분", "탐지할 댓글", "댓글 내용"]
    body_rows = list(ws2.iter_rows(min_row=2, values_only=True))
    assert len(body_rows) == 13  # 댓글 12개 + 게시글 링크 행
    assert body_rows[0][0] == "taboprou"  # 댓글1 계정
    assert body_rows[-1] == (None, None, None)  # 완료 링크 없음 → 빈 칸

    csv_dir = Path(export_out["comment_program_csv_dir"])
    csv_path = csv_dir / f"{KEYWORD}.csv"
    assert csv_path.exists()
    raw = csv_path.read_bytes()
    assert raw.startswith(b"\xef\xbb\xbf")  # UTF-8 BOM
    assert b"\r\n" in raw  # CRLF

    rows = list(rt.conn.execute("SELECT status FROM brand_queue WHERE brand='우아덤'"))
    assert rows and rows[0]["status"] == "exported"
    rt.close()


def test_placeholder_only_post_is_excluded_from_comment_program(tmp_path, monkeypatch):
    """댓글 12개가 전부 `테스트1`류 자리표시면 그 게시글은 댓글 프로그램용에서 뺀다."""
    rt = make_runtime(tmp_path)
    rt.settings.repo_root = tmp_path
    _patch_pool(monkeypatch, [{"keyword": KEYWORD, "cafe": "씨씨앙"}])
    monkeypatch.setattr(
        "v2r.content.bulk_generate.load_yaml",
        lambda name: {"bulk": {"allow_paid_fallback": False}} if name == "bulk" else {},
    )
    brand_queue.refill(rt, "우아덤", 1, target="vpc")
    # 검증이 자리표시 댓글을 걸러 재시도로 만들 수 있으니, 검증과 무관하게
    # 내보내기 로직만 확인하기 위해 수동으로 ready 상태 원고를 만든다
    out_dir = tmp_path / "warehouse" / "manuscripts" / "brand-queue" / "우아덤"
    out_dir.mkdir(parents=True, exist_ok=True)
    path = out_dir / f"{KEYWORD}.json"
    path.write_text(
        json.dumps(
            {
                "title": "t",
                "body": "b",
                "comments": [
                    {"label": lbl, "text": f"테스트{i or ''}"}
                    for i, lbl in enumerate(
                        ["댓글1", "대댓글1", "댓글2", "대댓글2", "대대댓글2", "대대대댓글2",
                         "댓글3", "대댓글3", "댓글4", "대댓글4", "댓글5", "대댓글5"]
                    )
                ],
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )
    row_id = list(rt.conn.execute("SELECT id FROM brand_queue WHERE brand='우아덤'"))[0]["id"]
    brand_queue.mark(rt, row_id, "ready", manuscript_path=str(path))

    export_out = bulk_generate.export_vpc(rt)
    assert KEYWORD in export_out["excluded"]
    assert export_out["exported"] == 0  # 자리표시 게시글이라 exported 처리 안 함
    rt.close()


def test_parser_routes_vpc_export_and_target_modifier():
    spec = parse_korean_command("가상pc 원고 내보내기")
    assert spec.task == "vpc_export"

    spec2 = parse_korean_command("우아덤 대량 원고 100건 가상pc")
    assert spec2.task == "bulk_generate" and spec2.target == "vpc"


def test_generate_all_uses_daily_quota_when_count_not_given(tmp_path, monkeypatch):
    rt = make_runtime(tmp_path)
    rt.settings.repo_root = tmp_path
    rt._llm = FakeLLM()
    rt._llm_ready = True
    _patch_pool(monkeypatch, [{"keyword": KEYWORD, "cafe": "씨씨앙"}])
    monkeypatch.setattr(
        "v2r.content.bulk_generate.load_yaml",
        lambda name: {
            "bulk": {
                "allow_paid_fallback": False,
                "brands": ["우아덤"],
                "daily_quota": {"우아덤": 1},
            }
        }
        if name == "bulk"
        else {},
    )
    out = bulk_generate.generate_all(rt, 0)
    assert out["brands"]["우아덤"]["generated"] == 1
    rt.close()


def test_schedule_bulk_entry_is_disabled():
    """대량 생성은 예약이 아니라 명령으로만 시작한다 (사용자 결정 2026-09-23)."""
    import yaml

    path = Path(__file__).resolve().parent.parent / "config" / "schedule.yaml"
    data = yaml.safe_load(path.read_text(encoding="utf-8"))
    entries = [e for e in data["entries"] if "대량 원고" in (e.get("command") or "")]
    assert entries, "대량 원고 예약 항목을 찾을 수 없습니다"
    assert all(e.get("enabled") is False for e in entries)


def test_parser_routes_bulk_generate():
    spec = parse_korean_command("우아덤 대량 원고 20건")
    assert spec.task == "bulk_generate" and spec.brand == "우아덤" and spec.count == 20

    spec_all = parse_korean_command("대량 원고 전체 20건")
    assert spec_all.task == "bulk_generate_all" and spec_all.count == 20

    spec_status = parse_korean_command("대량 원고 현황")
    assert spec_status.task == "bulk_generate_status"


def test_worker_dispatches_bulk_generate(tmp_path, monkeypatch):
    rt = make_runtime(tmp_path)
    rt.settings.repo_root = tmp_path
    rt._llm = FakeLLM()
    rt._llm_ready = True
    _patch_pool(monkeypatch, [{"keyword": KEYWORD, "cafe": "씨씨앙"}])
    spec = parse_korean_command("우아덤 대량 원고 1건")
    out = worker.handle_text(rt, "우아덤 대량 원고 1건")
    assert out["ok"] is True
    rt.close()
