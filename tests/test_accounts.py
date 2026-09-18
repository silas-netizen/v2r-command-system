"""계정 적재·규칙·배정 테스트."""

import pytest
from openpyxl import Workbook
from openpyxl.styles import PatternFill

from v2r.accounts.assign import AssignError, assign, rotate
from v2r.accounts.loader import Account, load_from_rows, load_from_xlsx
from v2r.accounts.rules import (
    CafeMatchError,
    eligible,
    find_affiliate,
    is_excluded,
    is_manager,
    work_type_for,
)

CAFES = {
    "affiliate": [
        {"name": "씨씨앙", "revision_delay_hours": 4, "aliases": ["ccang"]},
        {"name": "양평맘", "revision_delay_hours": 20, "aliases": []},
    ],
    "self_owned": [{"name": "고요한 아침", "cafe_id": 14567700}],
}


def _wb(path, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "아이디 리스트"
    ws.append(["A", "ID", "C", "D", "E", "F", "G", "작업 구분", "연동", "등급"])
    for r in rows:
        ws.append(r)
    return wb, ws


def test_load_from_xlsx_회색음영_제외(tmp_path):
    path = tmp_path / "accounts.xlsx"
    wb, ws = _wb(path, [
        ["", "alpha", "", "", "", "", "", "자사 카페", "V2R", ""],
        ["", "bravo", "", "", "", "", "", "제휴 작업", "V2R", ""],
        ["", "charlie", "", "", "", "", "", "자사 카페", "V2R", ""],
    ])
    ws.cell(row=4, column=2).fill = PatternFill(
        start_color="FFD9D9D9", end_color="FFD9D9D9", fill_type="solid"
    )
    wb.save(path)

    accounts, stats = load_from_xlsx(path, {"gray_rgb": ["D9D9D9"]})
    by_id = {a.login_id: a for a in accounts}
    assert stats["rows"] == 3
    assert stats["gray_excluded"] == 1
    assert stats["self_v2r"] == 1 and stats["affiliate_v2r"] == 1
    assert by_id["charlie"].excluded is True
    assert by_id["charlie"].shade == "회색"
    assert by_id["alpha"].excluded is False
    assert not hasattr(by_id["alpha"], "password")


def test_load_from_rows_별칭과_truthy():
    rows = [
        {"아이디": "a1", "작업구분": "자사 카페", "연동": "V2R", "등급": "", "닉네임": "닉"},
        {"ID": "a2", "작업 구분": "제휴 작업", "I": "V2R", "J": "m", "제외": "y"},
    ]
    accounts, stats = load_from_rows(rows)
    assert accounts[0].nickname == "닉" and accounts[0].account_type == "self_owned"
    assert accounts[1].excluded is True and stats["gray_excluded"] == 1


def test_rules_기본():
    mgr = Account("m1", "자사 카페", "V2R", grade="매니저")
    ok = Account("ok", "자사 카페", "V2R")
    assert is_manager(mgr) and not is_manager(ok)
    assert not is_excluded(ok)
    assert is_excluded(Account("g", "자사 카페", "V2R", excluded=True))


def test_work_type_for():
    assert work_type_for("publish_daily", "씨씨앙", CAFES) == "제휴 작업"
    assert work_type_for("publish_daily", "ccang", CAFES) == "제휴 작업"
    assert work_type_for("publish_brand", "", CAFES) == "제휴 작업"
    assert work_type_for("publish_info", "", CAFES) == "자사 카페"
    assert work_type_for("publish_daily", "고요한 아침", CAFES) == "자사 카페"


def test_카페_중복이름이면_에러():
    cfg = {"affiliate": [
        {"name": "씨씨앙", "aliases": []},
        {"name": "씨 씨 앙", "aliases": []},
    ]}
    with pytest.raises(CafeMatchError):
        find_affiliate("씨씨앙", cfg)


def test_eligible_필터():
    accounts = [
        Account("keep", "자사 카페", "V2R"),
        Account("gray", "자사 카페", "V2R", excluded=True),
        Account("mgr", "자사 카페", "V2R", grade="manager"),
        Account("other", "제휴 작업", "V2R"),
        Account("nolink", "자사 카페", ""),
        Account("cmt", "자사 카페", "V2R"),
        Account("rest", "자사 카페", "V2R"),
    ]
    out = eligible(accounts, "자사 카페", {"cmt"}, {"rest"})
    assert [a.login_id for a in out] == ["keep"]


def test_assign_manual_없는계정_에러():
    pool = [Account("a", "", ""), Account("b", "", "")]
    assert assign(pool, mode="manual", explicit=["b", "a"]) == ["b", "a"]
    with pytest.raises(AssignError) as exc:
        assign(pool, mode="manual", explicit=["a", "zzz"])
    assert "zzz" in str(exc.value)


def test_assign_auto_LRU_그리고_부족하면_에러():
    pool = ["a", "b", "c"]
    last = {"a": "2026-09-10T00:00:00", "b": None, "c": "2026-09-01T00:00:00"}
    assert assign(pool, mode="auto", count=2, last_used=last) == ["b", "c"]
    with pytest.raises(AssignError):
        assign(pool, mode="auto", count=4, last_used=last)  # 절대 채우지 않음
    with pytest.raises(AssignError):
        assign(pool, mode="auto", count=0, last_used=last)


def test_rotate():
    assert rotate(["a", "b"], 3) == "b"
    with pytest.raises(AssignError):
        rotate([], 0)
