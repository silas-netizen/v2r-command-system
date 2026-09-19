"""계정 시트 적재. 엑셀 회색 음영 / CSV 열 모두 지원. 비밀번호는 저장하지 않는다."""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

DEFAULT_SHEET = "아이디 리스트"
DEFAULT_GRAY = ["D9D9D9", "B7B7B7", "CCCCCC", "999999", "808080"]
DEFAULT_COLUMNS = {"id": "B", "work_type": "H", "linked": "I", "grade": "J"}

SELF_WORK_TYPE = "자사 카페"
AFFILIATE_WORK_TYPE = "제휴 작업"
#: 댓글 전용 계정의 작업 구분(시트 H열). 자사·제휴는 절대 섞어 쓰지 않는다(사용자 규칙 2026-09-19).
SELF_COMMENT_WORK_TYPE = "자사 댓글"
AFFILIATE_COMMENT_WORK_TYPE = "제휴 댓글"
LINKED_V2R = "V2R"

#: CSV 헤더 별칭
_ALIASES: dict[str, tuple[str, ...]] = {
    "login_id": ("ID", "id", "아이디", "B"),
    "work_type": ("작업 구분", "작업구분", "H"),
    "linked": ("연동", "I"),
    "grade": ("등급", "J"),
    "nickname": ("닉네임", "카페닉네임"),
    "shade": ("음영", "shade", "제외"),
}
_TRUTHY = {"y", "yes", "true", "1", "회색", "음영"}


class AccountLoadError(RuntimeError):
    """계정 적재 실패."""


@dataclass
class Account:
    """계정 한 건. 비밀번호 필드 없음."""

    login_id: str
    work_type: str = ""
    linked: str = ""
    excluded: bool = False
    grade: str = ""
    nickname: str = ""
    shade: str = ""
    account_type: str = ""


def _col_index(letter: str) -> int:
    """엑셀 열 문자 → 0-based 인덱스."""
    idx = 0
    for ch in letter.strip().upper():
        idx = idx * 26 + (ord(ch) - 64)
    return idx - 1


def _gray_set(cfg: dict | None) -> set[str]:
    raw = (cfg or {}).get("gray_rgb") or DEFAULT_GRAY
    return {str(x).strip().upper()[-6:] for x in raw}


def _is_gray(cell, grays: set[str]) -> bool:
    """셀 배경색이 회색 음영인지."""
    fill = getattr(cell, "fill", None)
    fg = getattr(fill, "fgColor", None)
    rgb = getattr(fg, "rgb", None)
    if not isinstance(rgb, str):
        return False
    tail = rgb.strip().upper()[-6:]
    return tail in grays


def _account_type(work_type: str) -> str:
    """work_type → affiliate/self_owned."""
    if work_type.strip() == AFFILIATE_WORK_TYPE:
        return "affiliate"
    if work_type.strip() == SELF_WORK_TYPE:
        return "self_owned"
    return ""


def load_from_xlsx(path: str | Path, cfg: dict | None = None) -> tuple[list[Account], dict]:
    """엑셀 계정 시트를 읽어 (계정 목록, 통계) 반환."""
    from openpyxl import load_workbook

    cfg = cfg or {}
    sheet_name = cfg.get("sheet_name") or DEFAULT_SHEET
    cols = {**DEFAULT_COLUMNS, **(cfg.get("columns") or {})}
    grays = _gray_set(cfg)

    wb = load_workbook(Path(path))
    if sheet_name not in wb.sheetnames:
        raise AccountLoadError(f"시트 없음: {sheet_name}")
    ws = wb[sheet_name]

    i_id = _col_index(cols["id"])
    i_wt = _col_index(cols["work_type"])
    i_ln = _col_index(cols["linked"])
    i_gr = _col_index(cols["grade"])

    accounts: list[Account] = []
    stats = {"rows": 0, "gray_excluded": 0, "self_v2r": 0, "affiliate_v2r": 0}

    for row in ws.iter_rows(min_row=2):
        if i_id >= len(row):
            continue
        cell = row[i_id]
        login_id = str(cell.value or "").strip()
        if not login_id:
            continue
        stats["rows"] += 1

        def val(i: int) -> str:
            return str(row[i].value or "").strip() if i < len(row) else ""

        gray = _is_gray(cell, grays)
        acc = Account(
            login_id=login_id,
            work_type=val(i_wt),
            linked=val(i_ln),
            excluded=gray,
            grade=val(i_gr),
            shade="회색" if gray else "",
        )
        acc.account_type = _account_type(acc.work_type)
        if gray:
            stats["gray_excluded"] += 1
        elif acc.linked.strip().upper() == LINKED_V2R:
            if acc.account_type == "self_owned":
                stats["self_v2r"] += 1
            elif acc.account_type == "affiliate":
                stats["affiliate_v2r"] += 1
        accounts.append(acc)

    wb.close()
    return accounts, stats


def _pick(row: dict, key: str) -> str:
    """별칭 표로 값 하나 꺼내기."""
    for alias in _ALIASES[key]:
        if alias in row and row[alias] is not None:
            return str(row[alias]).strip()
    lowered = {str(k).strip().casefold(): v for k, v in row.items()}
    for alias in _ALIASES[key]:
        v = lowered.get(alias.casefold())
        if v is not None:
            return str(v).strip()
    return ""


def load_from_rows(rows: list[dict]) -> tuple[list[Account], dict]:
    """CSV 행(dict) 목록 → (계정 목록, 통계)."""
    accounts: list[Account] = []
    stats = {"rows": 0, "gray_excluded": 0, "self_v2r": 0, "affiliate_v2r": 0}
    for row in rows or []:
        login_id = _pick(row, "login_id")
        if not login_id:
            continue
        stats["rows"] += 1
        shade_raw = _pick(row, "shade")
        excluded = shade_raw.strip().casefold() in _TRUTHY
        acc = Account(
            login_id=login_id,
            work_type=_pick(row, "work_type"),
            linked=_pick(row, "linked"),
            excluded=excluded,
            grade=_pick(row, "grade"),
            nickname=_pick(row, "nickname"),
            shade="회색" if excluded else "",
        )
        acc.account_type = _account_type(acc.work_type)
        if excluded:
            stats["gray_excluded"] += 1
        elif acc.linked.strip().upper() == LINKED_V2R:
            if acc.account_type == "self_owned":
                stats["self_v2r"] += 1
            elif acc.account_type == "affiliate":
                stats["affiliate_v2r"] += 1
        accounts.append(acc)
    return accounts, stats
